from __future__ import annotations

import io
import json
import logging
import re
import zipfile
from email import message_from_bytes
from email.policy import default as email_default_policy
from collections.abc import Mapping, Sequence
from typing import Any, assert_never

from graphon.file import file_manager
from graphon.file.enums import FileType
from graphon.file.models import File
from graphon.model_runtime.entities.message_entities import (
    AssistantPromptMessage,
    ImagePromptMessageContent,
    PromptMessage,
    PromptMessageContentType,
    PromptMessageContentUnionTypes,
    PromptMessageRole,
    SystemPromptMessage,
    TextPromptMessageContent,
    UserPromptMessage,
)
from graphon.model_runtime.entities.model_entities import (
    AIModelEntity,
    ModelPropertyKey,
)
from graphon.model_runtime.memory.prompt_message_memory import PromptMessageMemory
from graphon.nodes.base.entities import VariableSelector
from graphon.prompt_entities import MemoryConfig
from graphon.runtime.variable_pool import VariablePool
from graphon.template_rendering import Jinja2TemplateRenderer
from graphon.variables.segments import (
    ArrayAnySegment,
    ArrayFileSegment,
    FileSegment,
    NoneSegment,
)
from graphon.variables.template_resolution import convert_template

from .entities import (
    LLMNodeChatModelMessage,
    LLMNodeCompletionModelPromptTemplate,
)
from .exc import (
    InvalidVariableTypeError,
    MemoryRolePrefixRequiredError,
    NoPromptFoundError,
    TemplateTypeNotSupportError,
)
from .runtime_protocols import LLMProtocol

CONTEXT_PLACEHOLDER = "{{#context#}}"

logger = logging.getLogger(__name__)

VARIABLE_PATTERN = re.compile(r"\{\{#[^#]+#\}\}")
MAX_RESOLVED_VALUE_LENGTH = 1024


def fetch_model_schema(*, model_instance: LLMProtocol) -> AIModelEntity:
    model_schema = model_instance.get_model_schema()
    if not model_schema:
        msg = (
            "Model schema not found for "
            f"{getattr(model_instance, 'model_name', 'unknown model')}"
        )
        raise ValueError(msg)
    return model_schema


def build_model_identity_inputs(
    *,
    model_instance: LLMProtocol,
) -> dict[str, Any]:
    """Expose the prepared model identity in node inputs."""
    return {
        "model_provider": model_instance.provider,
        "model_name": model_instance.model_name,
    }


def fetch_files(variable_pool: VariablePool, selector: Sequence[str]) -> Sequence[File]:
    variable = variable_pool.get(selector)
    if variable is None:
        return []
    if isinstance(variable, FileSegment):
        return [variable.value]
    if isinstance(variable, ArrayFileSegment):
        return variable.value
    if isinstance(variable, NoneSegment | ArrayAnySegment):
        return []
    msg = f"Invalid variable type: {type(variable)}"
    raise InvalidVariableTypeError(msg)


def convert_history_messages_to_text(
    *,
    history_messages: Sequence[PromptMessage],
    human_prefix: str,
    ai_prefix: str,
) -> str:
    string_messages: list[str] = []
    for message in history_messages:
        if message.role == PromptMessageRole.USER:
            role = human_prefix
        elif message.role == PromptMessageRole.ASSISTANT:
            role = ai_prefix
        else:
            continue

        if isinstance(message.content, list):
            content_parts = []
            for content in message.content:
                if isinstance(content, TextPromptMessageContent):
                    content_parts.append(content.data)
                elif isinstance(content, ImagePromptMessageContent):
                    content_parts.append("[image]")

            inner_msg = "\n".join(content_parts)
            string_messages.append(f"{role}: {inner_msg}")
        else:
            string_messages.append(f"{role}: {message.content}")

    return "\n".join(string_messages)


def fetch_memory_text(
    *,
    memory: PromptMessageMemory,
    max_token_limit: int,
    message_limit: int | None = None,
    human_prefix: str = "Human",
    ai_prefix: str = "Assistant",
) -> str:
    history_messages = memory.get_history_prompt_messages(
        max_token_limit=max_token_limit,
        message_limit=message_limit,
    )
    return convert_history_messages_to_text(
        history_messages=history_messages,
        human_prefix=human_prefix,
        ai_prefix=ai_prefix,
    )


def _update_completion_prompt_content(
    *,
    prompt_messages: list[PromptMessage],
    memory_text: str,
    sys_query: str | None,
) -> None:
    prompt_content = prompt_messages[0].content
    if isinstance(prompt_content, str):
        prompt_messages[0].content = _prepend_or_replace_histories(
            prompt_content,
            memory_text,
        )
        if sys_query:
            prompt_messages[0].content = str(prompt_messages[0].content).replace(
                "#sys.query#",
                sys_query,
            )
        return

    if isinstance(prompt_content, list):
        _update_text_prompt_items(prompt_content, memory_text)
        if sys_query:
            _prepend_sys_query_to_text_prompt_items(prompt_content, sys_query)
        return

    msg = "Invalid prompt content type"
    raise ValueError(msg)


def _prepend_or_replace_histories(text: str, memory_text: str) -> str:
    if "#histories#" in text:
        return text.replace("#histories#", memory_text)
    return memory_text + "\n" + text


def _update_text_prompt_items(
    prompt_content: list[PromptMessageContentUnionTypes],
    memory_text: str,
) -> None:
    for content_item in prompt_content:
        if isinstance(content_item, TextPromptMessageContent):
            content_item.data = _prepend_or_replace_histories(
                content_item.data,
                memory_text,
            )


def _prepend_sys_query_to_text_prompt_items(
    prompt_content: list[PromptMessageContentUnionTypes],
    sys_query: str,
) -> None:
    for content_item in prompt_content:
        if isinstance(content_item, TextPromptMessageContent):
            content_item.data = sys_query + "\n" + content_item.data


def _filter_prompt_messages(
    *,
    prompt_messages: list[PromptMessage],
    model_schema: AIModelEntity,
) -> list[PromptMessage]:
    filtered_prompt_messages: list[PromptMessage] = []
    for prompt_message in prompt_messages:
        if isinstance(prompt_message.content, list):
            prompt_message_content: list[PromptMessageContentUnionTypes] = []
            for content_item in prompt_message.content:
                if not model_schema.supports_prompt_content_type(content_item.type):
                    continue
                prompt_message_content.append(content_item)
            if not prompt_message_content:
                continue
            if (
                len(prompt_message_content) == 1
                and prompt_message_content[0].type == PromptMessageContentType.TEXT
            ):
                prompt_message.content = prompt_message_content[0].data
            else:
                prompt_message.content = prompt_message_content
            filtered_prompt_messages.append(prompt_message)
        elif not prompt_message.is_empty():
            filtered_prompt_messages.append(prompt_message)

    if len(filtered_prompt_messages) == 0:
        msg = (
            "No prompt found in the LLM configuration. "
            "Please ensure a prompt is properly configured before proceeding."
        )
        raise NoPromptFoundError(msg)

    return filtered_prompt_messages


def fetch_prompt_messages(
    *,
    sys_query: str | None = None,
    sys_files: Sequence[File],
    context: str = "",
    memory: PromptMessageMemory | None = None,
    model_instance: LLMProtocol,
    prompt_template: Sequence[LLMNodeChatModelMessage]
    | LLMNodeCompletionModelPromptTemplate,
    stop: Sequence[str] | None = None,
    memory_config: MemoryConfig | None = None,
    vision_enabled: bool = False,
    vision_detail: ImagePromptMessageContent.DETAIL,
    variable_pool: VariablePool,
    jinja2_variables: Sequence[VariableSelector],
    context_files: list[File] | None = None,
    template_renderer: Jinja2TemplateRenderer | None = None,
) -> tuple[Sequence[PromptMessage], Sequence[str] | None]:
    prompt_messages: list[PromptMessage] = []
    model_schema = fetch_model_schema(model_instance=model_instance)

    match prompt_template:
        case list():
            prompt_messages.extend(
                handle_list_messages(
                    messages=prompt_template,
                    context=context,
                    jinja2_variables=jinja2_variables,
                    variable_pool=variable_pool,
                    vision_detail_config=vision_detail,
                    template_renderer=template_renderer,
                ),
            )
            prompt_messages.extend(
                handle_memory_chat_mode(
                    memory=memory,
                    memory_config=memory_config,
                    model_instance=model_instance,
                ),
            )

            if sys_query:
                prompt_messages.extend(
                    handle_list_messages(
                        messages=[
                            LLMNodeChatModelMessage(
                                text=sys_query,
                                role=PromptMessageRole.USER,
                                edition_type="basic",
                            ),
                        ],
                        context="",
                        jinja2_variables=[],
                        variable_pool=variable_pool,
                        vision_detail_config=vision_detail,
                        template_renderer=template_renderer,
                    ),
                )
        case LLMNodeCompletionModelPromptTemplate():
            prompt_messages.extend(
                handle_completion_template(
                    template=prompt_template,
                    context=context,
                    jinja2_variables=jinja2_variables,
                    variable_pool=variable_pool,
                    template_renderer=template_renderer,
                ),
            )

            memory_text = handle_memory_completion_mode(
                memory=memory,
                memory_config=memory_config,
                model_instance=model_instance,
            )
            _update_completion_prompt_content(
                prompt_messages=prompt_messages,
                memory_text=memory_text,
                sys_query=sys_query,
            )
        case _:
            raise TemplateTypeNotSupportError(type_name=str(type(prompt_template)))

    _append_file_prompts(
        prompt_messages=prompt_messages,
        files=sys_files,
        vision_enabled=vision_enabled,
        vision_detail=vision_detail,
    )
    _append_file_prompts(
        prompt_messages=prompt_messages,
        files=context_files or [],
        vision_enabled=vision_enabled,
        vision_detail=vision_detail,
    )
    return (
        _filter_prompt_messages(
            prompt_messages=prompt_messages,
            model_schema=model_schema,
        ),
        stop,
    )



_DOCUMENT_TEXT_EXTENSIONS = {
    "txt", "md", "mdx", "markdown", "csv", "html", "htm", "xml", "json", "eml", "svg",
}
_MAX_DOCUMENT_CHARS = 30000


def _uploaded_file_extension(file: File) -> str:
    return (file.extension or "").removeprefix(".").lower()


def _clip_document_text(text: str) -> str:
    if len(text) <= _MAX_DOCUMENT_CHARS:
        return text
    return text[:_MAX_DOCUMENT_CHARS] + "\n...[truncated]"


def _decode_uploaded_text(blob: bytes) -> str:
    if blob.startswith(b"\xff\xfe") or blob.startswith(b"\xfe\xff"):
        return blob.decode("utf-16", errors="replace")
    if blob.startswith(b"\xef\xbb\xbf"):
        return blob.decode("utf-8-sig", errors="replace")
    return blob.decode("utf-8", errors="replace")


def _strip_markup(value: str) -> str:
    value = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", value)
    value = re.sub(r"(?s)<[^>]+>", " ", value)
    return re.sub(r"\n{3,}", "\n\n", value).strip()


def _xlsx_text(blob: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if cell is None else str(cell) for cell in row]
                if any(value.strip() for value in values):
                    rows.append("\t".join(values))
            if rows:
                parts.append(f"[{sheet.title}]\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(parts)


def _xls_text(blob: bytes) -> str:
    import xlrd

    book = xlrd.open_workbook(file_contents=blob)
    parts: list[str] = []
    for sheet in book.sheets():
        rows: list[str] = []
        for row_index in range(sheet.nrows):
            values = [str(sheet.cell_value(row_index, column)) for column in range(sheet.ncols)]
            if any(value.strip() for value in values):
                rows.append("\t".join(values))
        if rows:
            parts.append(f"[{sheet.name}]\n" + "\n".join(rows))
    return "\n\n".join(parts)


def _docx_text(blob: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(blob))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                parts.append("\t".join(values))
    return "\n".join(parts)


def _pptx_text(blob: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(blob))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines = [shape.text.strip() for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
        if lines:
            parts.append(f"[slide {index}]\n" + "\n".join(lines))
    return "\n\n".join(parts)


def _pdf_text(blob: bytes) -> str:
    import pypdfium2

    document = pypdfium2.PdfDocument(blob)
    try:
        parts: list[str] = []
        for index in range(len(document)):
            page = document[index]
            text_page = page.get_textpage()
            parts.append(text_page.get_text_bounded())
            text_page.close()
            page.close()
        return "\n\n".join(parts)
    finally:
        document.close()


def _epub_text(blob: bytes) -> str:
    parts: list[str] = []
    with zipfile.ZipFile(io.BytesIO(blob)) as book:
        chapters = sorted(name for name in book.namelist() if name.lower().endswith((".xhtml", ".html", ".htm")))
        for chapter in chapters:
            text = _strip_markup(_decode_uploaded_text(book.read(chapter)))
            if text:
                parts.append(text)
    return "\n\n".join(parts)


def _eml_text(blob: bytes) -> str:
    message = message_from_bytes(blob, policy=email_default_policy)
    body = ""
    if message.is_multipart():
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                content = part.get_content()
                body = content if isinstance(content, str) else str(content)
                break
        if not body:
            for part in message.walk():
                if part.get_content_type() == "text/html":
                    content = part.get_content()
                    raw = content if isinstance(content, str) else str(content)
                    body = _strip_markup(raw)
                    break
    else:
        content = message.get_content()
        raw = content if isinstance(content, str) else str(content)
        body = _strip_markup(raw) if message.get_content_type() == "text/html" else raw
    header = "\n".join(
        f"{name}: {message.get(name, '')}" for name in ("From", "To", "Subject", "Date") if message.get(name)
    )
    return "\n".join(part for part in (header, body) if part)


def _read_uploaded_document(file: File) -> str:
    from graphon.file.runtime import get_workflow_file_runtime

    label = file.filename or "uploaded file"
    extension = _uploaded_file_extension(file)
    if extension in {"doc", "ppt", "msg"}:
        return f"[{label}] The file was uploaded, but .{extension} cannot be read by this chat node."
    try:
        blob = get_workflow_file_runtime().load_file_bytes(file=file)
    except Exception as error:
        return f"[{label}] The file was uploaded, but it could not be read: {error}"
    readers = {
        "xlsx": _xlsx_text,
        "xls": _xls_text,
        "docx": _docx_text,
        "pptx": _pptx_text,
        "pdf": _pdf_text,
        "epub": _epub_text,
        "eml": _eml_text,
    }
    try:
        if extension in readers:
            text = readers[extension](blob)
        elif extension == "json":
            decoded = _decode_uploaded_text(blob)
            try:
                text = json.dumps(json.loads(decoded), ensure_ascii=False, indent=2)
            except json.JSONDecodeError:
                text = decoded
        else:
            text = _decode_uploaded_text(blob)
            if extension in {"html", "htm", "xml", "svg"}:
                text = _strip_markup(text)
    except Exception as error:
        return f"[{label}] The file was uploaded, but it could not be read: {error}"
    text = text.strip()
    if not text:
        return f"[{label}] The file was uploaded, but no text could be read from it."
    return f"[{label}]\n{_clip_document_text(text)}"


def _prompt_contents_for_uploaded_file(
    file: File,
    image_detail_config: ImagePromptMessageContent.DETAIL,
) -> list[PromptMessageContentUnionTypes]:
    extension = _uploaded_file_extension(file)
    if extension == "svg" or file.type == FileType.DOCUMENT or extension in _DOCUMENT_TEXT_EXTENSIONS:
        return [TextPromptMessageContent(data=_read_uploaded_document(file))]
    if file.type in {FileType.IMAGE, FileType.VIDEO, FileType.AUDIO}:
        return [file_manager.to_prompt_message_content(file, image_detail_config=image_detail_config)]
    return []


def handle_list_messages(
    *,
    messages: Sequence[LLMNodeChatModelMessage],
    context: str,
    jinja2_variables: Sequence[VariableSelector],
    variable_pool: VariablePool,
    vision_detail_config: ImagePromptMessageContent.DETAIL,
    template_renderer: Jinja2TemplateRenderer | None = None,
) -> Sequence[PromptMessage]:
    prompt_messages: list[PromptMessage] = []
    for message in messages:
        if message.edition_type == "jinja2":
            result_text = render_jinja2_message(
                template=message.jinja2_text or "",
                jinja2_variables=jinja2_variables,
                variable_pool=variable_pool,
                template_renderer=template_renderer,
            )
            prompt_messages.append(
                combine_message_content_with_role(
                    contents=[TextPromptMessageContent(data=result_text)],
                    role=message.role,
                ),
            )
            continue

        template = message.text.replace(CONTEXT_PLACEHOLDER, context)
        segment_group = convert_template(variable_pool, template)
        file_contents: list[PromptMessageContentUnionTypes] = []
        for segment in segment_group.value:
            if isinstance(segment, ArrayFileSegment):
                for file in segment.value:
                    file_contents.extend(
                        _prompt_contents_for_uploaded_file(file, vision_detail_config)
                    )
            elif isinstance(segment, FileSegment):
                file_contents.extend(
                    _prompt_contents_for_uploaded_file(segment.value, vision_detail_config)
                )

        if segment_group.text:
            prompt_messages.append(
                combine_message_content_with_role(
                    contents=[TextPromptMessageContent(data=segment_group.text)],
                    role=message.role,
                ),
            )
        if file_contents:
            prompt_messages.append(
                combine_message_content_with_role(
                    contents=file_contents,
                    role=message.role,
                ),
            )

    return prompt_messages


def render_jinja2_message(
    *,
    template: str,
    jinja2_variables: Sequence[VariableSelector],
    variable_pool: VariablePool,
    template_renderer: Jinja2TemplateRenderer | None = None,
) -> str:
    if not template:
        return ""
    if template_renderer is None:
        msg = "template_renderer is required for jinja2 prompt rendering"
        raise ValueError(msg)

    jinja2_inputs: dict[str, Any] = {}
    for jinja2_variable in jinja2_variables:
        variable = variable_pool.get(jinja2_variable.value_selector)
        jinja2_inputs[jinja2_variable.variable] = (
            variable.to_object() if variable else ""
        )
    return template_renderer.render_template(template, jinja2_inputs)


def handle_completion_template(
    *,
    template: LLMNodeCompletionModelPromptTemplate,
    context: str,
    jinja2_variables: Sequence[VariableSelector],
    variable_pool: VariablePool,
    template_renderer: Jinja2TemplateRenderer | None = None,
) -> Sequence[PromptMessage]:
    if template.edition_type == "jinja2":
        result_text = render_jinja2_message(
            template=template.jinja2_text or "",
            jinja2_variables=jinja2_variables,
            variable_pool=variable_pool,
            template_renderer=template_renderer,
        )
    else:
        template_text = template.text.replace(CONTEXT_PLACEHOLDER, context)
        result_text = convert_template(variable_pool, template_text).text
    return [
        combine_message_content_with_role(
            contents=[TextPromptMessageContent(data=result_text)],
            role=PromptMessageRole.USER,
        ),
    ]


def combine_message_content_with_role(
    *,
    contents: str | list[PromptMessageContentUnionTypes] | None = None,
    role: PromptMessageRole,
) -> PromptMessage:
    match role:
        case PromptMessageRole.USER:
            return UserPromptMessage(content=contents)
        case PromptMessageRole.ASSISTANT:
            return AssistantPromptMessage(content=contents)
        case PromptMessageRole.SYSTEM:
            return SystemPromptMessage(content=contents)
        case PromptMessageRole.TOOL:
            msg = f"Role {role} is not supported"
            raise NotImplementedError(msg)
        case _:
            assert_never(role)


def calculate_rest_token(
    *,
    prompt_messages: list[PromptMessage],
    model_instance: LLMProtocol,
) -> int:
    rest_tokens = 2000
    runtime_model_schema = fetch_model_schema(model_instance=model_instance)
    runtime_model_parameters = model_instance.parameters

    model_context_tokens = runtime_model_schema.model_properties.get(
        ModelPropertyKey.CONTEXT_SIZE,
    )
    if model_context_tokens:
        curr_message_tokens = model_instance.get_llm_num_tokens(prompt_messages)

        max_tokens = 0
        for parameter_rule in runtime_model_schema.parameter_rules:
            if parameter_rule.name == "max_tokens" or (
                parameter_rule.use_template
                and parameter_rule.use_template == "max_tokens"
            ):
                max_tokens = (
                    runtime_model_parameters.get(parameter_rule.name)
                    or runtime_model_parameters.get(str(parameter_rule.use_template))
                    or 0
                )

        rest_tokens = model_context_tokens - max_tokens - curr_message_tokens
        rest_tokens = max(rest_tokens, 0)

    return rest_tokens


def handle_memory_chat_mode(
    *,
    memory: PromptMessageMemory | None,
    memory_config: MemoryConfig | None,
    model_instance: LLMProtocol,
) -> Sequence[PromptMessage]:
    if not memory or not memory_config:
        return []
    rest_tokens = calculate_rest_token(
        prompt_messages=[],
        model_instance=model_instance,
    )
    return memory.get_history_prompt_messages(
        max_token_limit=rest_tokens,
        message_limit=memory_config.window.size
        if memory_config.window.enabled
        else None,
    )


def handle_memory_completion_mode(
    *,
    memory: PromptMessageMemory | None,
    memory_config: MemoryConfig | None,
    model_instance: LLMProtocol,
) -> str:
    if not memory or not memory_config:
        return ""

    rest_tokens = calculate_rest_token(
        prompt_messages=[],
        model_instance=model_instance,
    )
    if not memory_config.role_prefix:
        msg = "Memory role prefix is required for completion model."
        raise MemoryRolePrefixRequiredError(msg)

    return fetch_memory_text(
        memory=memory,
        max_token_limit=rest_tokens,
        message_limit=memory_config.window.size
        if memory_config.window.enabled
        else None,
        human_prefix=memory_config.role_prefix.user,
        ai_prefix=memory_config.role_prefix.assistant,
    )


def _append_file_prompts(
    *,
    prompt_messages: list[PromptMessage],
    files: Sequence[File],
    vision_enabled: bool,
    vision_detail: ImagePromptMessageContent.DETAIL,
) -> None:
    if not vision_enabled or not files:
        return

    file_prompts = [
        file_manager.to_prompt_message_content(file, image_detail_config=vision_detail)
        for file in files
    ]
    if (
        prompt_messages
        and isinstance(prompt_messages[-1], UserPromptMessage)
        and isinstance(prompt_messages[-1].content, list)
    ):
        existing_contents = prompt_messages[-1].content
        prompt_messages[-1] = UserPromptMessage(
            content=file_prompts + existing_contents,
        )
    else:
        prompt_messages.append(UserPromptMessage(content=file_prompts))


def _coerce_resolved_value(raw: str) -> int | float | bool | str:
    """Try to restore the original type from a resolved template string.

    Variable references are always resolved to text, but completion params may
    expect numeric or boolean values (e.g. a variable that holds "0.7" mapped to
    the ``temperature`` parameter).  This helper attempts a JSON parse so that
    ``"0.7"`` → ``0.7``, ``"true"`` → ``True``, etc.  Plain strings that are not
    valid JSON literals are returned as-is.

    Returns:
        The parsed numeric or boolean literal, or the original string.

    """
    stripped = raw.strip()
    if not stripped:
        return raw

    try:
        parsed: Any = json.loads(stripped)
    except (json.JSONDecodeError, ValueError):
        return raw

    if isinstance(parsed, (int, float, bool)):
        return parsed
    return raw


def resolve_completion_params_variables(
    completion_params: Mapping[str, Any],
    variable_pool: VariablePool,
) -> dict[str, Any]:
    """Resolve variable references (``{{#node_id.var#}}``) in string-typed
    completion params.

    Security notes:
    - Resolved values are length-capped to ``MAX_RESOLVED_VALUE_LENGTH`` to
      prevent denial-of-service through excessively large variable payloads.
    - This follows the shared Graphon template-resolution path used across
      nodes (Answer Node, HTTP Request Node, Agent Node, etc.). The downstream
      model plugin receives these values as structured JSON key-value pairs — they
      are never concatenated into raw HTTP headers or SQL queries.
    - Numeric/boolean coercion is applied so that variables holding ``"0.7"`` are
      restored to their native type rather than sent as a bare string.

    Returns:
        Completion params with referenced variables resolved and coerced when possible.

    """
    resolved: dict[str, Any] = {}
    for key, value in completion_params.items():
        if isinstance(value, str) and VARIABLE_PATTERN.search(value):
            segment_group = convert_template(variable_pool, value)
            text = segment_group.text
            if len(text) > MAX_RESOLVED_VALUE_LENGTH:
                logger.warning(
                    "Resolved value for param '%s' truncated from %d to %d chars",
                    key,
                    len(text),
                    MAX_RESOLVED_VALUE_LENGTH,
                )
                text = text[:MAX_RESOLVED_VALUE_LENGTH]
            resolved[key] = _coerce_resolved_value(text)
        else:
            resolved[key] = value
    return resolved
